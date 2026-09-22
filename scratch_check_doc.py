import docx
import zipfile
import xml.etree.ElementTree as ET

def check_docx(filename):
    print("="*50)
    print("Checking:", filename)
    doc = docx.Document(filename)
    print(f"Number of sections: {len(doc.sections)}")
    for i, s in enumerate(doc.sections):
        sectPr = s._sectPr
        cols = sectPr.xpath('./w:cols')
        cols_info = ""
        if cols:
            num = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', 'default(1)')
            space = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space', 'none')
            cols_info = f"cols: num={num}, space={space}"
        else:
            cols_info = "cols: none"
        print(f"  Section {i}: start_type={s.start_type}, {cols_info}, top_margin={s.top_margin.pt}pt, bottom_margin={s.bottom_margin.pt}pt, left={s.left_margin.pt}pt, right={s.right_margin.pt}pt")
    
    print(f"Number of paragraphs: {len(doc.paragraphs)}")
    print(f"First 5 paragraphs:")
    for p in doc.paragraphs[:5]:
        print(f"  [{p.style.name}] {p.text[:60]}")
    print(f"Last 5 paragraphs:")
    for p in doc.paragraphs[-5:]:
        print(f"  [{p.style.name}] {p.text[:60]}")

try:
    check_docx("NEURALFLOW_COMPLETE_PROJECT_REPORT.docx")
except Exception as e:
    print("Error reading NEURALFLOW:", e)

try:
    # check template converted
    from report_builder_core import load_template_doc
    doc_tmpl = load_template_doc(r"c:\Users\Chaha\Downloads\Project-template-a4.docx")
    print("="*50)
    print("Checking original template:")
    print(f"Number of sections: {len(doc_tmpl.sections)}")
    for i, s in enumerate(doc_tmpl.sections):
        sectPr = s._sectPr
        cols = sectPr.xpath('./w:cols')
        cols_info = ""
        if cols:
            num = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', 'default(1)')
            space = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space', 'none')
            cols_info = f"cols: num={num}, space={space}"
        else:
            cols_info = "cols: none"
        print(f"  Section {i}: start_type={s.start_type}, {cols_info}, top_margin={s.top_margin.pt}pt, bottom_margin={s.bottom_margin.pt}pt, left={s.left_margin.pt}pt, right={s.right_margin.pt}pt")
    print(f"Number of paragraphs in template: {len(doc_tmpl.paragraphs)}")
    print("Paragraphs in template:")
    for i, p in enumerate(doc_tmpl.paragraphs):
        sect_in_p = p._p.xpath('./w:pPr/w:sectPr')
        sect_tag = f" [HAS sectPr: {len(sect_in_p)}]" if sect_in_p else ""
        if p.text.strip():
            print(f"  p{i} [{p.style.name}]{sect_tag}: {p.text[:60]}")
        elif sect_tag:
            print(f"  p{i} [EMPTY]{sect_tag}")
except Exception as e:
    print("Error reading template:", e)
