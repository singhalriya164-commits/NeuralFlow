import docx

doc = docx.Document(r"c:\Users\Chaha\Downloads\NeuralFlow_Group9_Final_Fixed.docx")
print("--- FIGURE CAPTION RUNS ---")
for i, p in enumerate(doc.paragraphs):
    if p.style.name == "figure caption":
        print(f"P{i}: {[r.text for r in p.runs]}")

print("\n--- TABLE HEAD RUNS ---")
for i, p in enumerate(doc.paragraphs):
    if p.style.name == "table head":
        print(f"P{i}: {[r.text for r in p.runs]}")
