import docx
from report_builder_core import load_template_doc
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

doc = load_template_doc(r"c:\Users\Chaha\Downloads\Project-template-a4.docx")

# Set title
p0 = doc.paragraphs[0]
p0.text = "NeuralFlow: A Controlled Empirical Benchmark, Gradient Dynamics Analysis, and Enterprise Platform for Recurrent Neural Network Architectures"

# Subtitle
p1 = doc.paragraphs[1]
p1.text = "*A Comprehensive Technical Whitepaper & Senior Project Research Report"

# p3 is the sectPr ending Section 0. Leave it alone.

# Authors in p4, p5, p6
authors = [
    "Chahat Deep Singh\nDept. of Computer Science & Engineering\nNeuralFlow AI Research Laboratory\nChandigarh, India\nchahat@neuralflow.io",
    "Aarav Sharma\nDept. of Information Technology\nSequence Modeling Group\nNew Delhi, India\naarav.sharma@research.ac.in",
    "Dr. Priya Venkatesh\nCenter for Computational Intelligence\nDept. of Artificial Intelligence\nBengaluru, India\np.venkatesh@univ.edu.in"
]

for idx, p_idx in enumerate([4, 5, 6]):
    p = doc.paragraphs[p_idx]
    p.text = authors[idx]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Clear p7
doc.paragraphs[7].text = ""

# Delete placeholder paragraphs from index 10 onwards
p_elements = [p._p for p in doc.paragraphs[10:]]
for p_elem in p_elements:
    p_elem.getparent().remove(p_elem)

for t in doc.tables:
    t._tbl.getparent().remove(t._tbl)

# Now configure the last section (the body) to be 2 columns
last_sec = doc.sections[-1]
sectPr = last_sec._sectPr

# Remove any existing cols
for c in sectPr.xpath('./w:cols'):
    sectPr.remove(c)

# Add 2 columns: 18pt space = 360 dxa
cols = parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="360"/>')
sectPr.append(cols)

# Ensure margins match template section 3:
# top=54pt (1080 dxa), bottom=72pt (1440 dxa), left=45.35pt (907 dxa), right=45.35pt (907 dxa)
for m in sectPr.xpath('./w:pgMar'):
    sectPr.remove(m)
pgMar = parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1080" w:right="907" w:bottom="1440" w:left="907" w:header="720" w:footer="720" w:gutter="0"/>')
sectPr.append(pgMar)

# Add Abstract
p_abs = doc.add_paragraph(style='Abstract')
r_lead = p_abs.add_run("Abstract—")
r_lead.font.bold = True
r_lead.font.italic = True
r_text = p_abs.add_run("This is the test abstract for NeuralFlow. It demonstrates strict adherence to the IEEE template. " * 5)

# Add Keywords
p_kw = doc.add_paragraph(style='Keywords')
r_kw_lead = p_kw.add_run("Keywords—")
r_kw_lead.font.bold = True
r_kw_lead.font.italic = True
r_kw_text = p_kw.add_run("Recurrent Neural Networks, LSTM, GRU, Gradient Dynamics, Spatial Trajectories.")

# Add Heading 1
p_h1 = doc.add_paragraph(style='Heading 1')
p_h1.add_run("I.  INTRODUCTION")

for i in range(10):
    p_body = doc.add_paragraph(style='Body Text')
    p_body.add_run(f"Paragraph {i+1}: Recurrent Neural Networks represent foundational sequential modeling tools. In this controlled benchmark, we evaluate gradient stability, accuracy, and computational efficiency across three distinct trajectory classes. " * 3)

doc.save("test_template_output.docx")
print("Saved test_template_output.docx successfully!")
