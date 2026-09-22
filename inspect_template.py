import zipfile, io, os
import docx
from docx.enum.style import WD_STYLE_TYPE

path = r'c:\Users\Chaha\Downloads\Project-template-a4.docx'

def s2t(c):
    for a, b in [
        (b'http://purl.oclc.org/ooxml/officeDocument/relationships', b'http://schemas.openxmlformats.org/officeDocument/2006/relationships'),
        (b'http://purl.oclc.org/ooxml/wordprocessingml/main', b'http://schemas.openxmlformats.org/wordprocessingml/2006/main'),
        (b'http://purl.oclc.org/ooxml/drawingml/main', b'http://schemas.openxmlformats.org/drawingml/2006/main'),
        (b'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing', b'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'),
        (b'http://purl.oclc.org/ooxml/officeDocument/math', b'http://schemas.openxmlformats.org/officeDocument/2006/math')
    ]:
        c = c.replace(a, b)
    return c

buf = io.BytesIO()
with zipfile.ZipFile(path, 'r') as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        zout.writestr(item, s2t(zin.read(item.filename)))
buf.seek(0)
doc = docx.Document(buf)

print(f"=== DOCUMENT SECTIONS ({len(doc.sections)}) ===")
for i, s in enumerate(doc.sections):
    sectPr = s._sectPr
    cols = sectPr.xpath('./w:cols')
    col_str = ""
    if cols:
        col_str = f"num={cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num')}, space={cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space')}"
    print(f"Section {i}: start={s.start_type}, w={s.page_width.pt}pt, h={s.page_height.pt}pt, margins(L={s.left_margin.pt}, R={s.right_margin.pt}, T={s.top_margin.pt}, B={s.bottom_margin.pt}), cols: {col_str}")
    for h in s.header.paragraphs:
        if h.text.strip(): print(f"  Header: {repr(h.text)}")
    for f in s.footer.paragraphs:
        if f.text.strip(): print(f"  Footer: {repr(f.text)}")

print(f"\n=== PARAGRAPH STYLES ===")
for s in doc.styles:
    if s.type == WD_STYLE_TYPE.PARAGRAPH:
        f_name = s.font.name if s.font else "inherited"
        f_sz = s.font.size.pt if (s.font and s.font.size) else "inherited"
        bld = s.font.bold if s.font else "inherited"
        it = s.font.italic if s.font else "inherited"
        print(f"Style: '{s.name}' -> font={f_name}, size={f_sz}, bold={bld}, italic={it}")

print(f"\n=== TABLE SAMPLE ===")
for i, t in enumerate(doc.tables):
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols, style='{t.style.name}'")
    for r_idx, r in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  Row {r_idx}: {cells}")
