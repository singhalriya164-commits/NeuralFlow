import docx

doc = docx.Document(r"c:\Users\Chaha\Downloads\NeuralFlow_Group9_Final_Fixed.docx")
print("Total paragraphs:", len(doc.paragraphs))

# Check for consecutive identical paragraphs
duplicates = []
for i in range(len(doc.paragraphs) - 1):
    t1 = doc.paragraphs[i].text.strip()
    t2 = doc.paragraphs[i+1].text.strip()
    if t1 and t1 == t2:
        duplicates.append((i, i+1, t1[:50]))

print(f"Consecutive identical paragraphs: {len(duplicates)}")
for d in duplicates:
    print(d)

# Check for duplicate headings
headings = {}
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith("Heading"):
        t = p.text.strip()
        if t in headings:
            print(f"Duplicate heading: {repr(t)} at P{headings[t]} and P{i}")
        else:
            headings[t] = i
