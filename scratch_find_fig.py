import docx
import re

doc = docx.Document(r"c:\Users\Chaha\Downloads\NeuralFlow_Group9_Final_Fixed.docx")

print("--- SEARCHING 'Fig' IN PARAGRAPHS ---")
for i, p in enumerate(doc.paragraphs):
    if re.search(r'\bfig\b|\bfig\.', p.text, re.IGNORECASE):
        print(f"P{i} [{p.style.name}]:")
        for match in re.finditer(r'(fig[a-z0-9\.\s\-,]+)', p.text, re.IGNORECASE):
            print(f"   Match: {repr(match.group(0))}")
        print(f"   Full text: {repr(p.text[:120])}")

print("\n--- SEARCHING 'Fig' IN TABLES ---")
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if re.search(r'\bfig\b|\bfig\.', cell.text, re.IGNORECASE):
                print(f"Table {t_idx} [R{r_idx}, C{c_idx}]: {repr(cell.text[:100])}")
