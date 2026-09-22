import docx
import re

doc = docx.Document(r"c:\Users\Chaha\Downloads\NeuralFlow_Group9_Final_Fixed.docx")
matches = []
for t_i, table in enumerate(doc.tables):
    for r_i, row in enumerate(table.rows):
        for c_i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                if re.search(r'\bfig\b|\bfig\.', p.text, re.I):
                    matches.append((t_i, r_i, c_i, p.text))
print("Table matches:", len(matches))
for m in matches:
    print(m)
