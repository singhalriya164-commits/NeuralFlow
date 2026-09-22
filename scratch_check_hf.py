import docx
import re

doc = docx.Document(r"c:\Users\Chaha\Downloads\NeuralFlow_Group9_Final_Fixed.docx")

print("Checking headers and footers:")
for i, s in enumerate(doc.sections):
    for p in s.header.paragraphs:
        if re.search(r'\bfig', p.text, re.I):
            print(f"Section {i} header: {p.text}")
    for p in s.footer.paragraphs:
        if re.search(r'\bfig', p.text, re.I):
            print(f"Section {i} footer: {p.text}")
            
print("Done checking headers/footers.")
